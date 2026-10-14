"""Step 7: build the 10-minute presentation (slides/presentation.pptx) from results and figures."""
import json

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

import config as C

AUTHOR = "José Alberto Martínez Navarro"
NAVY = RGBColor(0x1B, 0x2A, 0x49)
ACCENT = RGBColor(0xD6, 0x27, 0x28)
GREY = RGBColor(0x55, 0x55, 0x55)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_title(slide, text, sub=None):
    tb = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    p.text = text; p.font.size = Pt(30); p.font.bold = True; p.font.color.rgb = NAVY
    if sub:
        p2 = tb.text_frame.add_paragraph(); p2.text = sub; p2.font.size = Pt(16); p2.font.color.rgb = GREY
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.25), Inches(12.3), Emu(28000))
    line.fill.solid(); line.fill.fore_color.rgb = ACCENT; line.line.fill.background()


def add_bullets(slide, items, left=0.6, top=1.5, width=6.0, height=5.5, size=18):
    height = min(height, 7.3 - top)  # never extend past the slide bottom
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame; tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        level = 0
        if isinstance(it, tuple):
            it, level = it
        p.text = ("• " if level == 0 else "– ") + it
        p.level = level; p.font.size = Pt(size - 3 * level); p.space_after = Pt(6)


def add_image(slide, path, left, top, width=None, height=None):
    kw = {}
    if width: kw["width"] = Inches(width)
    if height: kw["height"] = Inches(height)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), **kw)


def add_table(slide, df, cols, headers, left, top, width, fmt, font=12):
    rows, ncols = len(df) + 1, len(cols)
    shape = slide.shapes.add_table(rows, ncols, Inches(left), Inches(top), Inches(width), Inches(0.4 * rows))
    t = shape.table
    for j, h in enumerate(headers):
        c = t.cell(0, j); c.text = h
        c.text_frame.paragraphs[0].font.size = Pt(font); c.text_frame.paragraphs[0].font.bold = True
    for i, (_, r) in enumerate(df.iterrows(), start=1):
        for j, col in enumerate(cols):
            v = r[col]
            c = t.cell(i, j); c.text = fmt.get(col, "{}").format(v) if pd.notna(v) else "–"
            c.text_frame.paragraphs[0].font.size = Pt(font)
            c.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER if j else PP_ALIGN.LEFT


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


df = pd.read_csv(C.RESULTS_DIR / "summary.csv")
stats = json.load(open(C.RESULTS_DIR / "dataset_stats.json"))
single = df[df.key != "ensemble"]
best = single.loc[single.macro_f1.idxmax()]
ens = df[df.key == "ensemble"].iloc[0] if (df.key == "ensemble").any() else None
F = C.FIGURES_DIR

# 1 Title -----------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK)
tb = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(2.5))
p = tb.text_frame.paragraphs[0]; p.text = "CNNs vs Vision Transformers for Skin-Lesion Classification"
p.font.size = Pt(36); p.font.bold = True; p.font.color.rgb = NAVY
p = tb.text_frame.add_paragraph(); p.text = "A controlled comparison of ResNet-50, EfficientNet-B0 and Swin-T on ISIC 2019"
p.font.size = Pt(22); p.font.color.rgb = GREY
p = tb.text_frame.add_paragraph(); p.text = f"\n{AUTHOR}  ·  Aprendizaje Profundo  ·  Final project, October 2026"; p.font.size = Pt(18)
notes(s, "Hello. My project compares three ImageNet-pretrained backbones on the ISIC 2019 dermoscopy dataset under one identical protocol, and looks at cost and interpretability, not only accuracy.")

# 2 Problem --------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Problem and objective")
add_bullets(s, [
    "Dermoscopic classification of 8 skin-lesion diagnoses (melanoma, nevus, BCC, AK, BKL, DF, vascular, SCC)",
    "Practical question: CNN or Transformer for a clinical prototype, and at what computational price?",
    "Published ISIC results are hard to compare: different splits, resolutions, ensembles, metadata",
    "Objective: hold everything fixed and compare ResNet-50, EfficientNet-B0 and Swin-T on",
    ("accuracy, balanced accuracy, macro-F1, AUC, per-class F1", 1),
    ("parameters, FLOPs, VRAM, training time, convergence, inference latency", 1),
    ("Grad-CAM evidence maps and a soft-voting ensemble (creativity)", 1),
], width=12)
notes(s, "The motivation is practical: a deployment decision depends on cost as much as on accuracy, and the literature rarely reports both under the same protocol.")

# 3 Dataset ---------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Dataset: ISIC 2019", f"{sum(stats['n'].values()):,} dermoscopic images · 8 classes · strong imbalance (NV ≈ 50 %)")
add_bullets(s, [
    f"Lesion-grouped, class-stratified split: {stats['n']['train']:,} train / {stats['n']['val']:,} val / {stats['n']['test']:,} test",
    "Same lesion never appears in two splits (lesion_id metadata) → no leakage",
    "Pre-resize to 256 px, train on random 224 crops + flips, rotation, colour jitter",
    "Eval on centre 224 crop; ImageNet normalisation",
], width=5.6, size=16)
add_image(s, F / "class_dist.png", 6.4, 1.6, width=6.6)
notes(s, "ISIC 2019 merges HAM10000, BCN20000 and MSK. I split by lesion id so that repeated images of the same lesion cannot leak from train to test, which inflates scores in many papers.")

# 4 Models & protocol -----------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Models and shared protocol")
add_bullets(s, [
    f"ResNet-50 (CNN, {single[single.key=='resnet50'].params_M.iloc[0]:.1f} M params)  ·  EfficientNet-B0 (efficient CNN, {single[single.key=='efficientnet_b0'].params_M.iloc[0]:.1f} M)  ·  Swin-T (hierarchical Transformer, {single[single.key=='swin_t'].params_M.iloc[0]:.1f} M)",
    "All ImageNet-1k pretrained (timm), new 8-way head, full fine-tuning",
    "Identical recipe: batch 32 · AdamW lr 1e-4, wd 0.05 · 1 warm-up + cosine · ≤12 epochs, early stop on val macro-F1",
    "Class-weighted cross-entropy (inverse √frequency) + label smoothing 0.1 · AMP fp16 · seed 42",
    "Single RTX 4070 Laptop GPU (8 GB) · best-val checkpoint evaluated once on test",
    "Swin-T chosen over ViT-B so that cost is comparable to ResNet-50 (fair comparison)",
], width=12, size=17)
notes(s, "Everything that is not the architecture is identical, so differences can be attributed to the backbone.")

# 5 Results table ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Results on the test split", "Accuracy and cost side by side")
cols = ["model", "accuracy", "balanced_accuracy", "macro_f1", "macro_auc", "params_M", "gflops", "latency_ms_b1", "train_time_min", "peak_vram_mb", "epochs_to_95"]
hdr = ["Model", "Acc", "Bal. acc", "Macro-F1", "AUC", "Params (M)", "GFLOPs", "Latency b1 (ms)", "Train (min)", "VRAM (MB)", "Ep. to 95 %"]
fmt = {"accuracy": "{:.3f}", "balanced_accuracy": "{:.3f}", "macro_f1": "{:.3f}", "macro_auc": "{:.3f}", "params_M": "{:.1f}",
       "gflops": "{:.2f}", "latency_ms_b1": "{:.1f}", "train_time_min": "{:.0f}", "peak_vram_mb": "{:.0f}", "epochs_to_95": "{:.0f}"}
add_table(s, df, cols, hdr, 0.5, 1.6, 12.3, fmt, font=13)
add_bullets(s, [f"Best single model: {best.model} ({100*best.macro_f1:.1f} % macro-F1)"] +
            ([f"Ensemble of 3: {100*ens.macro_f1:.1f} % macro-F1, AUC {ens.macro_auc:.3f}"] if ens is not None else []),
            top=4.3, width=12, size=18)
notes(s, "Read the table from left to right: the accuracy columns are close, the cost columns are not.")

# 6 Curves ---------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Convergence under the same schedule")
add_image(s, F / "curves.png", 0.5, 1.5, width=12.3)
add_bullets(s, [f"{r.model}: best epoch {int(r.best_epoch)}, 95 % of best F1 at epoch {int(r.epochs_to_95)}, {r.epoch_time_s:.0f} s/epoch" for _, r in single.iterrows()],
            top=5.2, width=12, size=15)
notes(s, "All three converge within the 12-epoch budget thanks to ImageNet pre-training; note which one gets there first and how the validation loss behaves.")

# 7 Cost -------------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Accuracy vs computational cost")
add_image(s, F / "acc_vs_cost.png", 0.5, 1.5, width=8.2)
add_bullets(s, [
    "x-axis: parameters / batch-1 fp16 latency; y-axis: test macro-F1",
    "Top-left is the sweet spot (accurate and cheap)",
    "Cost gap between models is much larger than the accuracy gap",
], left=8.9, top=1.8, width=4.2, size=15)
notes(s, "This is the figure I would show to an engineering manager: the choice is a budget decision.")

# 8 Per-class ------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Where the models differ: per-class F1 and confusions")
add_image(s, F / "per_class_f1.png", 0.4, 1.5, width=7.4)
add_image(s, F / f"cm_{best.key}.png", 8.0, 1.4, height=4.6)
add_bullets(s, ["Rare classes (DF, VASC, SCC) and MEL↔NV, AK↔BKL↔SCC confusions drive the macro-F1",
                f"Right: confusion matrix of {best.model}"], top=6.1, width=12, size=14)
notes(s, "Macro-F1 is dominated by the rare classes; class weighting helps but does not remove the MEL/NV and keratosis confusions that are also hard for dermatologists.")

# 9 Grad-CAM -------------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Interpretability: Grad-CAM (target = true class)")
add_image(s, F / "gradcam_grid.png", 2.2, 1.4, height=5.9)
notes(s, "The CNNs focus on compact lesion regions; Swin spreads attention over wider context. Failures often correspond to evidence on hair, ruler marks or borders.")

# 10 Ensemble / discussion -----------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Discussion")
items = [
    f"Accuracy: {best.model} is the best single backbone, but the spread across families is small",
    "Cost: EfficientNet-B0 is the cheapest to train and serve; Swin-T and ResNet-50 cost similar FLOPs but differ in latency and VRAM",
    "Generalisation: balanced accuracy and per-class F1 show the imbalance problem is only partly solved by loss weighting",
]
if ens is not None:
    items.append(f"Ensemble: {100*ens.macro_f1:.1f} % macro-F1 vs {100*best.macro_f1:.1f} % — the families are complementary, at 3× inference cost")
items += ["Pre-training matters more than architecture family on 25k images (consistent with Raghu et al. 2019, Matsoukas et al. 2021)"]
add_bullets(s, items, width=12, size=17)
notes(s, "Go beyond 'which is best': cost, interpretability and generalisation each tell a different story.")

# 11 Conclusions ---------------------------------------------------------------------------
s = prs.slides.add_slide(BLANK); add_title(s, "Conclusions and future work")
add_bullets(s, [
    "Under an identical, leakage-free protocol the CNN/Transformer gap on ISIC 2019 is small; the cost gap is large",
    "Recommendation: EfficientNet-B0 for edge/latency budgets, Swin-T or ResNet-50 when VRAM/latency allow, ensemble when accuracy is paramount",
    "Future work: higher resolution and multi-resolution ensembles (Gessert et al.), patient metadata, self-supervised pre-training on dermoscopy (DINO/MAE), test-time augmentation, calibration and external validation (official ISIC test set, PAD-UFES-20)",
    "Everything is reproducible: run_all.ps1 regenerates every number and figure",
], width=12, size=18)
notes(s, "Close with the recommendation and the reproducibility statement.")

out = C.PROJECT_DIR / "slides" / "presentation.pptx"
out.parent.mkdir(exist_ok=True)
prs.save(out)
print("saved", out)
